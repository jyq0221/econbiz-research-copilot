"""Check actual DOCX cells and record separately declared visual observations."""

import io

from .state import WorkflowError


def check_word_content(docx_bytes, display):
    try:
        from docx import Document
        doc = Document(io.BytesIO(docx_bytes))
    except ImportError as exc:
        raise WorkflowError('内容核对需要 documents 可选依赖') from exc
    except Exception as exc:
        raise WorkflowError('无法读取 Word 文档：'+str(exc)) from exc
    checks = [dict(name='table_count', passed=len(doc.tables) == len(display['tables']))]
    for i, expected in enumerate(display['tables']):
        actual = [] if i >= len(doc.tables) else [[cell.text for cell in row.cells] for row in doc.tables[i].rows]
        checks.append(dict(name=f'table_{i+1}', passed=actual == [expected['header']]+expected['rows']))
    paragraphs = [p.text for p in doc.paragraphs]
    expected_paragraphs = [display['title'], '本报告汇集已核验模型的比较结果、样本差异和变量说明，供研究讨论与修改使用。']
    for table in display['tables']:
        expected_paragraphs += [table['title']] + ['注：'+note for note in table['notes']]
    expected_paragraphs += ['样本与设定差异'] + display['paragraphs'] + ['来源与版本',
        '生成时间：'+display['created_at']+'。本文件为生成时快照，继续研究前请核对项目中的有效版本。'] + display['sources']
    checks.append(dict(name='paragraph_order_and_extras',
                       passed=[p for p in paragraphs if p] == [p for p in expected_paragraphs if p]))
    for text in [display['title']]+display['paragraphs']+display['sources']:
        checks.append(dict(name='paragraph', text=text, passed=text in paragraphs))
    for t in display['tables']:
        checks.append(dict(name='caption', passed=t['title'] in paragraphs))
        for note in t['notes']:
            checks.append(dict(name='note', passed='注：'+note in paragraphs))
    return dict(status='passed' if all(c['passed'] for c in checks) else 'failed', checks=checks,
                scope='实际表格和已登记文字；不认证视觉排版或因果解释')


def read_word_report(workspace, report_id):
    import json
    from .files import read_verified
    from .comparison import read_model_comparison
    from .comparison_display import build_display
    a = workspace.project.require_usable(report_id)
    if a['kind'] != 'word_report':
        raise WorkflowError('需要 Word 交付记录')
    try:
        content = a['content']
        records = [f for f in a['files'] if f['path'].endswith('/delivery.json')]
        docs = [f for f in a['files'] if f['path'].endswith('.docx')]
        if len(records) != 1 or len(docs) != 1 or json.loads(read_verified(workspace.root,records[0])) != content:
            raise WorkflowError('Word 记录与登记材料不一致')
        comparison = read_model_comparison(workspace, content['comparison']['id'])
        if comparison['version'] != content['comparison']['version']:
            raise WorkflowError('比较版本已变更')
        display = build_display(comparison['content'])
        if content['narrative'] is not None:
            ref = content['narrative']
            note = workspace.project.require_usable(ref['id'])
            if note['version'] != ref['version'] or note['kind'] != 'session_note':
                raise WorkflowError('解释笔记已变更')
            display['paragraphs'].append('研究说明（来源 '+ref['id']+'）：'+note['content']['summary'])
        display['sources'].append(f'比较 {comparison["id"]} v{comparison["version"]}')
        if display != content['display'] or docs[0]['sha256'] != content['docx_sha256']:
            raise WorkflowError('Word 展示内容与当前来源不一致')
        checked = check_word_content(read_verified(workspace.root,docs[0]), display)
        if checked['status'] != 'passed' or checked != content['numeric_text_check']:
            raise WorkflowError('Word 实际内容核对未通过')
        return a
    except (KeyError, ValueError, TypeError) as exc:
        raise WorkflowError('Word 记录损坏：'+str(exc)) from exc


def record_document_review(workspace, review_id, report_id, *, render_files, inspected_pages, findings, reason):
    from pathlib import Path
    from .state import now, require_text
    from .workspace import encode_json
    from .files import describe_file, read_verified
    report = read_word_report(workspace,report_id)
    if (not isinstance(render_files,dict) or set(render_files) != {'docx_sha256','pdf','pages'} or
            render_files['docx_sha256'] != report['content']['docx_sha256']):
        raise WorkflowError('渲染材料须绑定当前 DOCX 散列')
    if not isinstance(findings,list):
        raise WorkflowError('视觉发现须为文本列表')
    for finding in findings: require_text(finding,'视觉发现')
    try:
        from pypdf import PdfReader
        from PIL import Image
    except ImportError as exc:
        raise WorkflowError('版式记录需要 document-review 可选依赖') from exc

    def external(path):
        path = Path(path)
        ref = describe_file(workspace.root,path.absolute(),scope='external',role='render_evidence')
        return read_verified(workspace.root,ref)
    pdf = external(render_files['pdf'])
    try:
        reader = PdfReader(io.BytesIO(pdf), strict=True)
        count = len(reader.pages)
    except Exception as exc:
        raise WorkflowError('渲染 PDF 无法解析：'+str(exc)) from exc
    if (not count or not isinstance(inspected_pages,list) or
            any(type(n) is not int for n in inspected_pages) or sorted(inspected_pages) != list(range(1,count+1))):
        raise WorkflowError('须逐页检查全部实际 PDF 页面')
    pages = render_files['pages']
    if not isinstance(pages,list) or len(pages) != count or len(set(pages)) != count:
        raise WorkflowError('逐页图片数量与实际 PDF 不一致')
    images = []
    for i,path in enumerate(pages,1):
        if Path(path).name != f'page-{i}.png':
            raise WorkflowError('页图片须按 page-1.png 起连续排序')
        raw = external(path)
        try:
            with Image.open(io.BytesIO(raw)) as image:
                if image.format != 'PNG': raise ValueError('not PNG')
                image.verify()
        except Exception as exc:
            raise WorkflowError('页图片无效：'+str(exc)) from exc
        images.append(raw)
    version = workspace._version_number(review_id)
    prefix = f'{workspace.layout["research_reports"]}/{review_id}/v{version:04d}'
    content = dict(report={'id':report_id,'version':report['version']},
        docx_sha256=report['content']['docx_sha256'], recorded_at=now(),
        render_check='passed', visual_check='failed' if findings else 'passed',
        inspected_pages=inspected_pages, page_count=count, findings=findings,
        scope='渲染文件结构与散列已检查；逐页视觉观察由调用者声明，不认证观察者身份或像素对应关系')
    writes = [(prefix+'/render.pdf',pdf)]+[(prefix+f'/page-{i}.png',raw) for i,raw in enumerate(images,1)]
    writes.append((prefix+'/review.json',encode_json(content)))
    files = [workspace._file(p,b,'document_review') for p,b in writes]
    staged = workspace._stage(review_id,'document_review',content,[report_id],files,reason)
    staged.mark(review_id,'completed','passed','检查证据已保存；视觉结论单独记录')
    workspace._publish(staged,writes)
    return workspace.project.artifact(review_id)


def read_document_review(workspace, review_id):
    import json
    from .files import read_verified
    a = workspace.project.require_usable(review_id)
    if a['kind'] != 'document_review':
        raise WorkflowError('需要版式检查记录')
    try:
        c = a['content']
        refs = [f for f in a['files'] if f['path'].endswith('/review.json')]
        if len(refs) != 1 or json.loads(read_verified(workspace.root, refs[0])) != c:
            raise WorkflowError('版式检查记录与登记材料不一致')
        report = read_word_report(workspace, c['report']['id'])
        if report['version'] != c['report']['version'] or report['content']['docx_sha256'] != c['docx_sha256']:
            raise WorkflowError('版式检查对应的 Word 已变更')
        return a
    except (KeyError, TypeError, ValueError) as exc:
        raise WorkflowError('版式检查记录损坏：'+str(exc)) from exc
