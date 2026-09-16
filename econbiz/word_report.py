"""Optional editable Word delivery consuming a current model comparison."""

import io

from .comparison import read_model_comparison
from .comparison_display import build_display
from .state import WorkflowError, now
from .workspace import encode_json


def _docx_api():
    try:
        from docx import Document
        from docx.shared import Mm, Pt
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.enum.section import WD_SECTION, WD_ORIENT
        return Document, Mm, Pt, OxmlElement, qn, WD_SECTION, WD_ORIENT
    except ImportError as exc:
        raise WorkflowError('Word 导出需要 documents 可选依赖；现有比较记录仍可使用') from exc


def render_word(display):
    Document, Mm, Pt, Element, qn, Sections, Orient = _docx_api()
    doc = Document()
    doc.core_properties.title = display['title']
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    for name in ('Normal','Title','Heading 1','Heading 2'):
        style = doc.styles[name]
        style.font.name = 'Arial'
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
            fonts.attrib.pop(qn('w:'+key), None)
        fonts.set(qn('w:eastAsia'), 'SimSun')
        for border in style.element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
        style.font.color.rgb = __import__('docx.shared', fromlist=['RGBColor']).RGBColor(0,0,0)
    doc.styles['Normal'].font.size = Pt(10.5)
    doc.styles['Normal'].paragraph_format.space_after = Pt(5)
    doc.styles['Title'].font.size = Pt(18)
    doc.styles['Heading 1'].font.size = Pt(13)

    def layout(section, landscape):
        section.orientation = Orient.LANDSCAPE if landscape else Orient.PORTRAIT
        section.page_width, section.page_height = (Mm(297),Mm(210)) if landscape else (Mm(210),Mm(297))
        section.top_margin = section.bottom_margin = Mm(18)
        section.left_margin = section.right_margin = Mm(18)
    landscape = bool(display['tables'] and display['tables'][0]['kind'] == 'models' and len(display['tables'][0]['header']) > 4)
    layout(doc.sections[0], landscape)
    doc.add_paragraph(display['title'], 'Title')
    doc.add_paragraph('本报告汇集已核验模型的比较结果、样本差异和变量说明，供研究讨论与修改使用。')
    for item in display['tables']:
        wide = item['kind'] == 'models' and len(item['header']) > 4
        if wide != landscape:
            layout(doc.add_section(Sections.NEW_PAGE), wide)
            landscape = wide
        heading = doc.add_paragraph(item['title'], 'Heading 1')
        heading.paragraph_format.keep_with_next = True
        if wide:
            heading.paragraph_format.space_before = Pt(10)
        table = doc.add_table(rows=1, cols=len(item['header']))
        table.autofit = False
        usable = 261 if wide else 174
        first = 40 if item['kind'] == 'models' else 32
        widths = [first]+[(usable-first)/(len(item['header'])-1)]*(len(item['header'])-1)
        if item['kind'] == 'definitions':
            widths = [35, 50, 15, 25, 49]
        for col, width in zip(table.columns, widths): col.width = Mm(width)
        borders = Element('w:tblBorders')
        for edge in ('top','bottom','left','right','insideH','insideV'):
            e = Element('w:'+edge)
            e.set(qn('w:val'), 'single' if edge in ('top','bottom') else 'nil')
            e.set(qn('w:sz'), '8')
            borders.append(e)
        table._tbl.tblPr.append(borders)
        repeat = Element('w:tblHeader')
        table.rows[0]._tr.get_or_add_trPr().append(repeat)
        all_rows = [item['header']]+item['rows']
        for row_index, values in enumerate(all_rows):
            row = table.rows[0] if row_index == 0 else table.add_row()
            row._tr.get_or_add_trPr().append(Element('w:cantSplit'))
            for i,(cell,value) in enumerate(zip(row.cells,values)):
                cell.width = Mm(widths[i])
                cell.text = str(value)
                for para in cell.paragraphs:
                    para.paragraph_format.keep_with_next = row_index == 0 or (len(all_rows) <= 8 and row_index < len(all_rows)-1)
                    para.paragraph_format.space_after = Pt(1.5 if wide else 3)
                    para.paragraph_format.space_before = Pt(1.5 if wide else 3)
                    for run in para.runs:
                        run.font.size = Pt(9)
                        run.bold = row_index == 0
                if row_index == 0:
                    cb = Element('w:tcBorders')
                    edge = Element('w:bottom'); edge.set(qn('w:val'),'single'); edge.set(qn('w:sz'),'6')
                    cb.append(edge); cell._tc.get_or_add_tcPr().append(cb)
        for note in item['notes']:
            doc.add_paragraph('注：'+note)
    if landscape:
        layout(doc.add_section(Sections.NEW_PAGE), False)
    doc.add_paragraph('样本与设定差异', 'Heading 1')
    for paragraph in display['paragraphs']: doc.add_paragraph(paragraph)
    doc.add_paragraph('来源与版本', 'Heading 1')
    doc.add_paragraph('生成时间：'+display['created_at']+'。本文件为生成时快照，继续研究前请核对项目中的有效版本。')
    for source in display['sources']: doc.add_paragraph(source)
    stream = io.BytesIO(); doc.save(stream)
    return stream.getvalue()


def write_word_report(workspace, report_id, comparison_id, *, narrative_id=None):
    comparison = read_model_comparison(workspace, comparison_id)
    display = build_display(comparison['content'])
    dependencies = [comparison_id]
    narrative = None
    if narrative_id is not None:
        narrative = workspace.project.require_usable(narrative_id)
        if narrative['kind'] != 'session_note':
            raise WorkflowError('结果说明须引用有效会话研究笔记')
        display['paragraphs'].append('研究说明（来源 '+narrative_id+'）：'+narrative['content']['summary'])
        dependencies.append(narrative_id)
    display['sources'].append(f'比较 {comparison_id} v{comparison["version"]}')
    raw = render_word(display)
    from .document_checks import check_word_content
    checked = check_word_content(raw, display)
    if checked['status'] != 'passed':
        raise WorkflowError('Word 内容核对未通过，未发布新文档')
    version = workspace._version_number(report_id)
    prefix = f'{workspace.layout["research_reports"]}/{report_id}/v{version:04d}'
    word_path = prefix+'/report.docx'
    ref = workspace._file(word_path, raw, 'word_report')
    content = dict(comparison={'id':comparison_id,'version':comparison['version']}, generated_at=now(),
        generated=True, numeric_text_check=checked, render_check='not_performed', visual_check='not_performed',
        docx_sha256=ref['sha256'], narrative=None if narrative is None else
        {'id':narrative_id,'version':narrative['version']}, display=display)
    writes = [(word_path,raw),(prefix+'/delivery.json',encode_json(content))]
    files = [ref, workspace._file(writes[1][0],writes[1][1],'word_delivery_record')]
    staged = workspace._stage(report_id,'word_report',content,dependencies,files,'生成可编辑模型比较 Word')
    staged.mark(report_id,'completed','passed','文件和数值文本已核对；渲染及视觉状态另行记录')
    workspace._publish(staged,writes)
    return workspace.project.artifact(report_id)
