"""CSV/XLSX adapters preserving source bytes, text identifiers and missing labels."""

import csv
import io
import math
import re
from pathlib import Path

from .execution_spec import read_csv
from .files import read_verified
from .state import WorkflowError


def import_table(workspace, artifact_id, source_id, *, sheet=None):
    source = workspace.project.require_usable(source_id)
    if source['kind'] != 'source' or len(source.get('files', [])) != 1:
        raise WorkflowError('表格接入需要单个已登记来源')
    ref = source['files'][0]
    raw = read_verified(workspace.root, ref)
    suffix = Path(ref['path']).suffix.lower()
    conversion = dict(format=suffix, source_sha256=ref['sha256'], source_id=source_id,
                      source_version=source['version'], sheet=sheet)
    if suffix == '.csv':
        if sheet is not None:
            raise WorkflowError('CSV 没有工作表选项')
        read_csv(raw)
        converted = raw
        conversion['rule'] = '原字节保留；不转换代码或缺失标记'
    elif suffix == '.xlsx':
        converted, metadata = _xlsx(raw, sheet)
        conversion.update(metadata)
    else:
        raise WorkflowError('首版表格适配支持 .csv/.xlsx；请保留原件并另存受支持格式')
    version = workspace._version_number(artifact_id)
    relative = f'{workspace.layout["data_interim"]}/{artifact_id}/v{version:04d}/table.csv'
    files = [workspace._file(relative, converted, 'analysis_table')]
    content = dict(role='raw_data', conversion=conversion)
    staged = workspace._stage(artifact_id, 'source', content, [source_id], files, '表格适配，保留原始来源')
    # Check actual registered bytes again before publishing any derived version.
    workspace.project.require_usable(source_id)
    staged.mark(artifact_id, 'completed', 'passed', '仅验证格式转换，尚需盘点与测量核对')
    workspace._publish(staged, [(relative, converted)])
    return workspace.project.artifact(artifact_id)


def _xlsx(raw, sheet):
    if not isinstance(sheet, str) or not sheet.strip():
        raise WorkflowError('XLSX 须显式指定工作表名称')
    try:
        import openpyxl
    except ImportError as exc:
        raise WorkflowError('Excel 接入需要安装分析依赖：python3 -m pip install ".[analysis]"') from exc
    try:
        book = openpyxl.load_workbook(io.BytesIO(raw), data_only=False, read_only=False)
    except Exception as exc:
        raise WorkflowError(f'XLSX 无法读取：{exc}') from exc
    try:
        if sheet not in book.sheetnames:
            raise WorkflowError(f'工作表不存在：{sheet}；现有工作表：{book.sheetnames}')
        ws = book[sheet]
        if ws.merged_cells.ranges:
            raise WorkflowError('工作表含合并单元格，须先明确整理表头和数据')
        rows, types = [], {}
        for cells in ws.iter_rows():
            row = []
            for cell in cells:
                value = cell.value
                if cell.data_type in {'f', 'e'}:
                    raise WorkflowError(f'{cell.coordinate} 包含公式或 Excel 错误值；须另存明确的原始值')
                if value is None:
                    text = ''
                elif isinstance(value, str):
                    text = value
                elif type(value) in {int, float} and math.isfinite(value):
                    if re.search(r'0{2,}', cell.number_format.split('.')[0]):
                        raise WorkflowError(f'{cell.coordinate} 的数字格式可能补足代码前导零；须先转为明确文本')
                    text = str(value)
                else:
                    raise WorkflowError(f'{cell.coordinate} 存在日期、布尔或非有限值；请明确转换语义')
                types[cell.data_type] = types.get(cell.data_type, 0) + 1
                row.append(text)
            rows.append(row)
        # Trailing genuinely empty spreadsheet rows have no observations.
        while rows and all(value == '' for value in rows[-1]):
            rows.pop()
        stream = io.StringIO(newline='')
        csv.writer(stream, lineterminator='\n').writerows(rows)
        output = stream.getvalue().encode('utf-8')
        read_csv(output)
        return output, dict(sheet=sheet, rows=len(rows) - 1, columns=rows[0], cell_types=types,
                            rule='保留文本、空值和缺失词；数值使用存储值；公式/日期/编码补零拒绝猜测',
                            adapter_version=openpyxl.__version__)
    finally:
        book.close()
