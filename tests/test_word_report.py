import importlib.util
import io
import tempfile
import unittest
from pathlib import Path
from test_execution import HAS_ANALYSIS, study
from test_comparison import comparison_args
from econbiz.state import WorkflowError

HAS_DOCX = importlib.util.find_spec('docx') is not None

@unittest.skipUnless(HAS_ANALYSIS and HAS_DOCX, 'analysis/documents extra not installed')
class WordReportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.w = study(Path(self.temp.name))
        from econbiz.execution import execute_analysis
        from econbiz.comparison import write_model_comparison
        execute_analysis(self.w, 'r1', 'p1')
        execute_analysis(self.w, 'r2', 'p1')
        self.c = write_model_comparison(self.w, 'cmp', **comparison_args())

    def api(self):
        self.assertIsNotNone(importlib.util.find_spec('econbiz.word_report'))
        from econbiz.word_report import write_word_report
        return write_word_report

    def test_editable_table_numeric_text_and_status(self):
        from docx import Document
        a = self.api()(self.w, 'word', 'cmp')
        ref = next(f for f in a['files'] if f['path'].endswith('.docx'))
        doc = Document(io.BytesIO((self.w.root/ref['path']).read_bytes()))
        cells = [cell.text for t in doc.tables for row in t.rows for cell in row.cells]
        self.assertIn('观测数', cells)
        self.assertIn('72', cells)
        self.assertEqual(a['content']['numeric_text_check']['status'], 'passed')
        self.assertEqual(a['content']['visual_check'], 'not_performed')
        self.assertIn('tblHeader', doc.tables[0].rows[0]._tr.xml)
        self.assertNotIn('<w:drawing', doc._element.xml)
        self.w.project.mark('r1', 'completed', 'failed', 'test')
        with self.assertRaises(WorkflowError): self.w.project.require_usable('word')

    def test_generation_failure_keeps_previous_report(self):
        from unittest.mock import patch
        write = self.api()
        a = write(self.w, 'word', 'cmp')
        with patch('econbiz.word_report.render_word', side_effect=WorkflowError('missing dependency')):
            with self.assertRaises(WorkflowError): write(self.w, 'word', 'cmp')
        self.assertEqual(self.w.project.artifact('word'), a)

    def test_title_has_no_template_border_or_theme_font_override(self):
        from docx import Document
        from econbiz.word_report import render_word
        from econbiz.comparison_display import build_display
        doc=Document(io.BytesIO(render_word(build_display(self.c['content']))))
        self.assertNotIn('pBdr',doc.styles['Title'].element.xml)
        self.assertNotIn('eastAsiaTheme',doc.styles['Title'].element.xml)

    def test_resume_includes_delivery_status_and_paths(self):
        from econbiz.progress import resume_context
        from econbiz.workspace import Workspace
        self.api()(self.w,'word','cmp')
        summary=resume_context(Workspace.open(self.w.root))
        self.assertIn('deliveries',summary)
        self.assertEqual(summary['deliveries']['word']['content']['visual_check'],'not_performed')
        self.assertTrue(any(f['path'].endswith('.docx') for f in summary['deliveries']['word']['files']))

    def test_resume_rejects_changed_comparison_content(self):
        from econbiz.progress import resume_context
        from econbiz.workspace import Workspace
        self.api()(self.w,'word','cmp')
        self.w.project._state['artifacts']['cmp']['content']['models'][0]['parameters'][0]['coefficient']=999
        self.w.save()
        summary=resume_context(Workspace.open(self.w.root))
        self.assertFalse(summary['deliveries']['cmp']['usable'])
        self.assertFalse(summary['deliveries']['word']['usable'])
