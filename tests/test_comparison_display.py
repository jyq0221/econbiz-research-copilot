import importlib.util
import unittest
from econbiz.state import WorkflowError

class DisplayTests(unittest.TestCase):
    def test_unrounded_p_strict_boundary(self):
        self.assertIsNotNone(importlib.util.find_spec('econbiz.comparison_display'))
        from econbiz.comparison_display import significance_marks
        rules = [(0.01,'***'),(0.05,'**'),(0.1,'*')]
        self.assertEqual(significance_marks(.049999, rules), '**')
        self.assertEqual(significance_marks(.05, rules), '*')
        self.assertEqual(significance_marks(.1, rules), '')
        self.assertEqual(significance_marks(.001, []), '')
        with self.assertRaises(WorkflowError): significance_marks(.1, [(float('nan'),'*')])

    def test_zero_missing_and_precision(self):
        self.assertIsNotNone(importlib.util.find_spec('econbiz.comparison_display'))
        from econbiz.comparison_display import format_number
        self.assertEqual(format_number(0), '0.000')
        self.assertEqual(format_number(-.00001), '0.000')
        self.assertEqual(format_number(None), '未提供')
        with self.assertRaises(WorkflowError): format_number(float('inf'))

    @unittest.skipUnless(importlib.util.find_spec('docx'),'documents extra not installed')
    def test_wide_table_starts_landscape_without_empty_cover(self):
        from econbiz.word_report import render_word
        from docx import Document
        from docx.enum.section import WD_ORIENT
        import io
        d=dict(title='宽表',created_at='2026-09-16',paragraphs=[],sources=[],tables=[
            dict(title='模型',header=['变量']+[str(i) for i in range(6)],
                 rows=[['x']+['1.234']*6],notes=[],kind='models')])
        doc=Document(io.BytesIO(render_word(d)))
        self.assertEqual(doc.sections[0].orientation,WD_ORIENT.LANDSCAPE)
        self.assertTrue(doc.tables[0].rows[0].cells[0].paragraphs[0].paragraph_format.keep_with_next)

    def test_sample_rule_is_readable(self):
        from econbiz.comparison_display import describe_sample_rule
        self.assertEqual(describe_sample_rule({'field':'year','op':'ge','value':2020}), 'year 不小于 2020')
        self.assertEqual(describe_sample_rule({'complete_case':'c'}),'剔除 c 缺失的观测')
